from docx import Document
import sys
import glob

def extract_text(doc_path):
    doc = Document(doc_path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text.append(row_text)
    return "\n".join(text)

files = glob.glob("*.docx")
for f in files:
    try:
        content = extract_text(f)
        out_name = f.replace(".docx", ".txt")
        with open(out_name, "w", encoding="utf-8") as out:
            out.write(content)
        print(f"Extracted {f} to {out_name}")
    except Exception as e:
        print(f"Error on {f}: {e}")
